from PIL import Image
from pathlib import Path
import pytesseract
import os
import argparse
import shutil
from docx import Document
from docx.shared import Inches

if shutil.which("tesseract") is None:
    print("Tesseract not found in PATH")
    exit()

parser = argparse.ArgumentParser(description="Text from images in a folder to word files")
parser.add_argument("--include-images", action="store_true")
# parser.add_argument("--psm", default=6, help="Tesseract page segmentation mode")
args = parser.parse_args()

def get_input_path():
    dangerous_dir = [
    Path("/"),
    Path("/System"),
    Path("C:/"),
    Path("C:/Users"),
    Path("C:/Windows"),
    Path("C:/Program Files"),
    Path(os.path.expanduser("~"))]
    c = input("Enter full path of the folder:").strip()
    file_path = Path(c)

    if file_path.exists():
        print(f"The path '{file_path}' exists.")
    else:
        print(f"The path '{file_path}' does not exist.")
        exit()
    if file_path in dangerous_dir:
        print("Warning:System Directory Exiting.......")
        exit()
    return file_path

def write(path):
    document = Document()
    document.add_heading(path.name, 0)
    for a in path.rglob("*"):
        if a.suffix.lower() in [".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"]:
            try:
                p = pytesseract.image_to_string(Image.open(str(a)))
                document.add_heading(a.name, level=1)
                document.add_paragraph(p.strip())
                if args.include_images:
                    document.add_picture(str(a),width=Inches(3.0))
            except Exception as e:
                print(f"Error processing {a}: {e}")
    new_path = path.parent / f"{path.stem}.docx"
    document.save(new_path)


c = get_input_path()
write(c)




