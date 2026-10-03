from PIL import Image
from pathlib import Path
import pytesseract
from pytesseract import Output
import os
import argparse
import shutil
import json
import requests
from docx import Document
from docx.shared import Inches


IMAGE_EXTENSIONS = [".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"]
JEV_URL = "https://api.typesafe.ai/v1/systemone"
DEFAULT_JEV_MODEL = "jev-1.13.0"

BLOCK_TYPES = {
    "title": "A document title or main heading that identifies the document.",
    "heading": "A section heading or subsection heading.",
    "paragraph": "Normal continuous prose forming a paragraph.",
    "list": "A list item or a group of short list-like entries.",
    "caption": "A caption describing an image, figure, table, or other visual element.",
    "metadata": "Document metadata such as author, date, address, reference number, or similar information.",
    "unknown": "Does not clearly fit another category."
}


def get_input_path(path):
    dangerous_dir = [
        Path("/"),
        Path("/System"),
        Path("C:/"),
        Path("C:/Users"),
        Path("C:/Windows"),
        Path("C:/Program Files"),
        Path(os.path.expanduser("~")),
    ]

    file_path = Path(path)

    if not file_path.exists():
        print("The path '{}' does not exist.".format(file_path))
        raise SystemExit(1)

    if file_path in dangerous_dir:
        print("Warning: system directory. Exiting.......")
        raise SystemExit(1)

    if file_path.is_file():
        if file_path.suffix.lower() not in IMAGE_EXTENSIONS:
            print("The file is not a supported image.")
            raise SystemExit(1)

    elif not file_path.is_dir():
        print("Path is neither an image nor a directory.")
        raise SystemExit(1)

    return file_path


def ocr_image(image_path):
    """
    Run Tesseract and preserve layout information.

    Each returned block contains text plus bounding-box and OCR-confidence
    information. This is the main architectural change from Wordex v0.
    """
    image = Image.open(str(image_path))

    data = pytesseract.image_to_data(
        image,
        output_type=Output.DICT
    )

    words = []

    for i, text in enumerate(data["text"]):
        text = text.strip()

        if not text:
            continue

        try:
            confidence = float(data["conf"][i])
        except (ValueError, TypeError):
            confidence = -1.0

        words.append({
            "text": text,
            "left": int(data["left"][i]),
            "top": int(data["top"][i]),
            "width": int(data["width"][i]),
            "height": int(data["height"][i]),
            "confidence": confidence,
            "block_num": int(data["block_num"][i]),
            "par_num": int(data["par_num"][i]),
            "line_num": int(data["line_num"][i]),
        })

    # Group words into paragraph-level OCR blocks.
    grouped = {}

    for word in words:
        key = (
            word["block_num"],
            word["par_num"],
        )
        grouped.setdefault(key, []).append(word)

    blocks = []

    for block_id, block_words in grouped.items():
        block_words.sort(key=lambda w: (w["top"], w["left"]))

        lines = {}
        for word in block_words:
            lines.setdefault(word["line_num"], []).append(word)

        line_texts = []
        for line_num in sorted(lines):
            line_words = sorted(lines[line_num], key=lambda w: w["left"])
            line_texts.append(" ".join(w["text"] for w in line_words))

        left = min(w["left"] for w in block_words)
        top = min(w["top"] for w in block_words)
        right = max(w["left"] + w["width"] for w in block_words)
        bottom = max(w["top"] + w["height"] for w in block_words)

        confidences = [
            w["confidence"] for w in block_words
            if w["confidence"] >= 0
        ]

        blocks.append({
            "text": "\n".join(line_texts).strip(),
            "bbox": {
                "left": left,
                "top": top,
                "right": right,
                "bottom": bottom,
                "width": right - left,
                "height": bottom - top,
            },
            "ocr_confidence": (
                sum(confidences) / len(confidences)
                if confidences else -1.0
            ),
            "word_count": len(block_words),
            "block_num": block_id[0],
            "paragraph_num": block_id[1],
        })

    blocks.sort(key=lambda b: (b["bbox"]["top"], b["bbox"]["left"]))

    return {
        "image_width": image.width,
        "image_height": image.height,
        "blocks": blocks,
    }


def build_jev_state(image_name, ocr_result):
    """
    Keep Jev's state compact. It gets the information needed for a structural
    decision, rather than the full raw Tesseract word table.
    """
    state = {
        "image": image_name,
        "image_size": {
            "width": ocr_result["image_width"],
            "height": ocr_result["image_height"],
        },
        "blocks": [],
    }

    for index, block in enumerate(ocr_result["blocks"]):
        state["blocks"].append({
            "id": index,
            "text": block["text"],
            "bbox": block["bbox"],
            "ocr_confidence": round(block["ocr_confidence"], 1),
            "word_count": block["word_count"],
        })

    return state


def classify_blocks_with_jev(image_name, ocr_result, api_key, model):
    """
    Ask Jev to classify every OCR block in one request.

    Jev is used only for semantic decisions. Python remains responsible for
    rendering the resulting Word document.
    """
    blocks = ocr_result["blocks"]

    if not blocks:
        return []

    state = build_jev_state(image_name, ocr_result)

    questions = {}

    for index in range(len(blocks)):
        questions["block_{}".format(index)] = {
            "type": "choice",
            "instructions": (
                "What structural role does OCR block {} play in this document? "
                "Use the block text and its position/size relative to the page. "
                "Choose exactly one category."
            ).format(index),
            "criteria": BLOCK_TYPES,
        }

    response = requests.post(
        JEV_URL,
        headers={
            "Authorization": "Bearer " + api_key,
            "Content-Type": "application/json",
        },
        json={
            "model": model,
            "state": state,
            "questions": questions,
        },
        timeout=30,
    )

    if response.status_code != 200:
        raise RuntimeError(
            "Jev API returned HTTP {}: {}".format(
                response.status_code,
                response.text
            )
        )

    payload = response.json()
    answers = payload.get("answers", {})

    decisions = []

    for index, block in enumerate(blocks):
        answer = answers.get("block_{}".format(index), {})

        decisions.append({
            "id": index,
            "text": block["text"],
            "role": answer.get("choice", "paragraph"),
            "confidence": answer.get("confidence", 0.0),
            "probabilities": answer.get("probabilities", {}),
            "ocr_confidence": block["ocr_confidence"],
            "bbox": block["bbox"],
        })

    return decisions


def render_block(document, decision):
    """
    Deterministic Word rendering.

    Jev decides WHAT the block is. Python decides HOW that role is rendered.
    """
    text = decision["text"]
    role = decision["role"]

    if role == "title":
        document.add_heading(text.replace("\n", " "), level=0)

    elif role == "heading":
        document.add_heading(text.replace("\n", " "), level=1)

    elif role == "list":
        for line in text.splitlines():
            line = line.strip()
            if line:
                document.add_paragraph(line, style="List Bullet")

    elif role == "caption":
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text.replace("\n", " "))
        run.italic = True

    elif role == "metadata":
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text.replace("\n", " "))
        run.bold = True

    else:
        document.add_paragraph(text)


def write_document(path, api_key, model, include_images):
    document = Document()
    document.add_heading(path.name, 0)

    debug = []

    if path.is_file():
        image_paths = [path]
    else:
        image_paths = path.rglob("*")

    for image_path in image_paths:
        if image_path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue

        try:
            print("OCR: {}".format(image_path.name))
            ocr_result = ocr_image(image_path)

            print(
                "  Found {} OCR blocks. Asking Jev...".format(
                    len(ocr_result["blocks"])
                )
            )

            decisions = classify_blocks_with_jev(
                image_path.name,
                ocr_result,
                api_key,
                model,
            )

            document.add_heading(image_path.name, level=1)

            for decision in decisions:
                render_block(document, decision)

            if include_images:
                document.add_picture(
                    str(image_path),
                    width=Inches(3.0)
                )

            debug.append({
                "image": image_path.name,
                "blocks": decisions,
            })

        except Exception as e:
            print("Error processing {}: {}".format(image_path, e))

    new_path = path.parent / "{}_v1.docx".format(path.stem)
    document.save(str(new_path))

    debug_path = path.parent / "{}_v1_debug.json".format(path.stem)
    with open(str(debug_path), "w") as f:
        json.dump(debug, f, indent=2)

    print("")
    print("Created: {}".format(new_path))
    print("Debug decisions: {}".format(debug_path))


def main():
    if shutil.which("tesseract") is None:
        print("Tesseract not found in PATH")
        raise SystemExit(1)

    api_key = os.environ.get("TYPESAFE_API_KEY")

    if not api_key:
        print("TYPESAFE_API_KEY is not set.")
        raise SystemExit(1)

    parser = argparse.ArgumentParser(
        description="Wordex v1: OCR + layout-aware Jev classification + DOCX"
    )

    parser.add_argument("path", help="Path to an image or folder of images")

    parser.add_argument(
        "--include-images",
        action="store_true",
        help="Include the source images in the output document.",
    )
    parser.add_argument(
        "--model",
        default=os.environ.get("TYPESAFE_MODEL", DEFAULT_JEV_MODEL),
        help="Jev model to use.",
    )

    args = parser.parse_args()

    input_path = get_input_path(args.path)

    write_document(
        input_path,
        api_key,
        args.model,
        args.include_images,
    )


if __name__ == "__main__":
    main()
